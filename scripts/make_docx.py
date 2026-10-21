#!/usr/bin/env python3
"""Builds public/downloads/Quant-Overledger-Redbelly-Report.docx from REPORT.md (pandoc), then
applies the brand fonts and colours and clears the author properties. Run after make_pdf.py."""
import re, subprocess, tempfile
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R = Path(__file__).resolve().parent.parent
out = R / "public" / "downloads" / "Quant-Overledger-Redbelly-Report.docx"
out.parent.mkdir(parents=True, exist_ok=True)
md = (R / "REPORT.md").read_text()
title = re.search(r"^# (.+)$", md, re.M)[1]

tmp = Path(tempfile.mkdtemp()) / "raw.docx"
subprocess.run(["pandoc", str(R / "REPORT.md"), "-f", "markdown+pipe_tables", "-o", str(tmp)], check=True)

d = Document(str(tmp))
FONT, STEEL = "Be Vietnam Pro", RGBColor(0x1D, 0x31, 0x40)
for st in d.styles:
    if st.type != 1 or not hasattr(st, "font"): continue
    st.font.name = FONT
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if st.name.startswith(("Heading", "Title")): st.font.color.rgb = STEEL
    if st.name in ("Normal", "Body Text", "First Paragraph", "Compact"): st.font.size = Pt(10.5)
    if st.name in ("Source Code", "Verbatim Char"):
        st.font.name = "JetBrains Mono"; st.font.size = Pt(9)

# tables: thin borders and a steel header row
for t in d.tables:
    tblPr = t._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), "D8E0E5"); b.append(x)
    tblPr.append(b)
    for c in t.rows[0].cells:
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "1D3140"); c._tc.get_or_add_tcPr().append(sh)
        for p in c.paragraphs:
            for r in p.runs: r.font.color.rgb = RGBColor(255, 255, 255); r.font.bold = True
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(9)

# no author, no personal properties
cp = d.core_properties
cp.title = title; cp.author = ""; cp.last_modified_by = ""; cp.comments = ""; cp.keywords = ""; cp.subject = ""
d.save(str(out)); print("docx ok")
